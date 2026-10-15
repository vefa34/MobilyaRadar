"""CRM-ready CSV/XLSX export with licence-aware field filtering.

A field is exported only if at least one *exportable* source provides it (evidence rows).
E.g. a phone number known only from Google Places is withheld, because Places content may not be
exported in bulk. Synthetic demo records are excluded unless explicitly requested and are then
labelled. Companies marked do-not-contact are exported without contact fields.
"""
from __future__ import annotations

import csv
import io
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Company, CompanyContact, CompanyEvidence, OrganizationLead, SourceRegistry
from app.services.scoring import priority_label

COLUMNS = [
    ("company_id", "Company ID"), ("company_name", "Company Name"), ("legal_name", "Legal Name"), ("country", "Country"),
    ("city", "City"), ("address", "Address"), ("website", "Website"), ("general_email", "Email"), ("phone", "Phone"),
    ("purchasing_contact", "Purchasing Contact"), ("purchasing_email", "Purchasing Email"),
    ("primary_buyer_type", "Buyer Type"), ("buyer_types", "All Buyer Types"), ("product_categories", "Products"),
    ("import_status", "Import Activity"), ("store_count", "Store Count"), ("buyer_fit_score", "Buyer Fit Score"),
    ("data_confidence_score", "Data Confidence"), ("priority", "Priority"), ("registration_number", "Registration No"),
    ("vat_number", "VAT No"), ("research_status", "Research Status"), ("crm_status", "CRM Status"), ("sources", "Sources"),
    ("evidence_urls", "Evidence URLs"), ("last_verified_at", "Last Verified"), ("is_synthetic", "Synthetic (demo)"),
]
FIELD_EVIDENCE = {"company_name": "name", "address": "address", "website": "website", "general_email": "email", "phone": "phone",
                  "city": "city", "registration_number": "registration_number", "vat_number": "vat_number"}


def build_rows(db: Session, org_id, *, company_ids: list | None = None, job_id=None, min_fit: float | None = None,
               include_synthetic: bool = False, include_rejected: bool = False) -> tuple[list[dict], dict]:
    exportable = {k: v for k, v in db.execute(select(SourceRegistry.key, SourceRegistry.export_allowed))}
    q = select(Company).where(Company.org_id == org_id)
    if company_ids:
        q = q.where(Company.id.in_(company_ids))
    if job_id:
        q = q.where(Company.id.in_(select(CompanyEvidence.company_id).where(CompanyEvidence.job_id == job_id)))
    if min_fit is not None:
        q = q.where(Company.buyer_fit_score >= min_fit)
    if not include_synthetic:
        q = q.where(Company.is_synthetic.is_(False))
    if not include_rejected:
        q = q.where(Company.research_status != "rejected")
    companies = db.scalars(q.order_by(Company.buyer_fit_score.desc().nullslast())).all()
    stats = {"exported": 0, "withheld_fields": 0, "excluded_companies": 0}
    rows = []
    for c in companies:
        ev = db.scalars(select(CompanyEvidence).where(CompanyEvidence.company_id == c.id)).all()
        allowed_fields: dict[str, set] = defaultdict(set)
        urls = []
        for e in ev:
            if exportable.get(e.source_key, True):
                if e.field and e.value:
                    allowed_fields[e.field].add(e.value.lower())
                if e.source_url and e.source_url not in urls and not e.source_url.startswith("upload://"):
                    urls.append(e.source_url)
        if not allowed_fields.get("name"):
            stats["excluded_companies"] += 1  # name only known from non-exportable sources
            continue
        lead = db.scalar(select(OrganizationLead).where(OrganizationLead.org_id == org_id, OrganizationLead.company_id == c.id))
        dnc = lead is not None and lead.crm_status == "do_not_contact"
        purch = db.get(CompanyContact, c.purchasing_contact_id) if c.purchasing_contact_id else None
        row = {
            "company_id": str(c.id), "company_name": c.display_name, "legal_name": c.legal_name, "country": c.country, "city": c.city,
            "address": c.address, "website": c.website, "general_email": c.general_email, "phone": c.phone,
            "purchasing_contact": (purch.person_name or purch.job_title) if purch and exportable.get(purch.source_key or "", True) else None,
            "purchasing_email": purch.value if purch and exportable.get(purch.source_key or "", True) else None,
            "primary_buyer_type": c.primary_buyer_type, "buyer_types": ", ".join(c.buyer_types or []),
            "product_categories": ", ".join(c.product_categories or []), "import_status": c.import_activity_status,
            "store_count": c.store_count, "buyer_fit_score": c.buyer_fit_score, "data_confidence_score": c.data_confidence_score,
            "priority": priority_label(c.buyer_fit_score or 0, c.data_confidence_score or 0),
            "registration_number": c.registration_number, "vat_number": c.vat_number, "research_status": c.research_status,
            "crm_status": lead.crm_status if lead else "new", "sources": ", ".join(c.sources or []), "evidence_urls": " | ".join(urls[:5]),
            "last_verified_at": c.last_verified_at.isoformat() if c.last_verified_at else None, "is_synthetic": c.is_synthetic,
        }
        for col, field in FIELD_EVIDENCE.items():
            val = row.get(col)
            if col == "company_name" or val in (None, ""):
                continue
            if str(val).lower() not in allowed_fields.get(field, set()) and not _domain_ok(field, val, allowed_fields):
                row[col] = None
                stats["withheld_fields"] += 1
        if dnc:
            for col in ("general_email", "phone", "purchasing_contact", "purchasing_email"):
                row[col] = None
        rows.append(row)
    stats["exported"] = len(rows)
    return rows, stats


def _domain_ok(field: str, val, allowed: dict) -> bool:
    if field == "website":
        v = str(val).lower().rstrip("/")
        return any(a.rstrip("/") == v for a in allowed.get("website", set()))
    return False


def to_csv(rows: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([label for _, label in COLUMNS])
    for r in rows:
        w.writerow(["" if r.get(k) is None else r.get(k) for k, _ in COLUMNS])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 (Arabic/Georgian) correctly


def to_xlsx(rows: list[dict], meta: dict | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append([label for _, label in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
    for r in rows:
        ws.append(["" if r.get(k) is None else r.get(k) for k, _ in COLUMNS])
    ws.freeze_panes = "B2"
    for col in ws.columns:
        width = min(50, max(10, max(len(str(c.value or "")) for c in col[:200]) + 2))
        ws.column_dimensions[col[0].column_letter].width = width
    if meta:
        m = wb.create_sheet("Export info")
        for k, v in meta.items():
            m.append([k, str(v)])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
