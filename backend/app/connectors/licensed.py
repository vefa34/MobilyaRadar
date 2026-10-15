"""Skeletons for licensed / contract-only sources and the user-import connector.

Licensed connectors (Panjiva, ImportGenius, Volza, Kompass, Europages, wlw, Apollo, BIA, …) never
pretend to work. Without a confirmed licence they raise ``AccessRestrictedError`` so the planner
marks the cell ``access_restricted`` and the coverage report lists the gap. With a confirmed licence
and credentials they still refuse until a provider-specific adapter is implemented against the
provider's actual API documentation (subclass and override ``discover``/``enrich``).
"""
from __future__ import annotations

import csv
import io

from app.connectors.base import (
    AccessRestrictedError,
    ConnectorContext,
    DiscoveryQuery,
    DiscoveryResult,
    HealthStatus,
    LeadSourceConnector,
    RawCompany,
)


class LicensedSourceConnector(LeadSourceConnector):
    def __init__(self, source_key: str, *a, license_confirmed: bool = False, **kw):
        super().__init__(*a, **kw)
        self.key = source_key
        self.license_confirmed = license_confirmed

    def _refuse(self):
        if not self.license_confirmed:
            raise AccessRestrictedError(f"{self.key}: licence/contract not confirmed for this organisation")
        if not self.credentials:
            raise AccessRestrictedError(f"{self.key}: licence confirmed but no credentials stored")
        raise AccessRestrictedError(f"{self.key}: provider adapter not implemented yet — requires provider API documentation")

    async def discover(self, query: DiscoveryQuery, context: ConnectorContext) -> DiscoveryResult:
        self._refuse()
        raise AssertionError  # pragma: no cover

    async def enrich(self, company, context):
        self._refuse()

    async def health_check(self) -> HealthStatus:
        if not self.license_confirmed:
            return HealthStatus("pending_license", "licence/contract not confirmed")
        return HealthStatus("pending_license", "licence confirmed; provider adapter pending implementation")


HEADER_MAP = {
    "name": {"name", "company", "company name", "firma", "firmenname", "raison sociale", "société", "şirket", "şirket adı", "nom"},
    "website": {"website", "web", "url", "site", "homepage", "web sitesi"},
    "country": {"country", "land", "pays", "ülke", "country code"},
    "city": {"city", "stadt", "ville", "şehir", "town", "locality"},
    "address": {"address", "adresse", "adres", "street"},
    "phone": {"phone", "telephone", "telefon", "tel", "téléphone"},
    "email": {"email", "e-mail", "mail", "e-posta"},
    "categories": {"category", "categories", "kategorie", "catégorie", "kategori", "type", "activity"},
    "registration_number": {"registration number", "reg no", "company number", "handelsregister", "siren", "ico"},
    "vat_number": {"vat", "vat number", "ust-idnr", "tva", "vergi no"},
}


def parse_tabular(content: bytes, filename: str) -> list[dict]:
    """Parse CSV or XLSX into list of dicts with canonical keys."""
    rows: list[dict] = []
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        headers = [str(h or "").strip() for h in next(it, [])]
        for r in it:
            rows.append({headers[i]: ("" if v is None else str(v)) for i, v in enumerate(r) if i < len(headers)})
    else:
        text = content.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t") if text.strip() else csv.excel
        rows = list(csv.DictReader(io.StringIO(text), dialect=dialect))
    out = []
    for r in rows:
        canon: dict = {}
        for k, v in r.items():
            key = (k or "").strip().lower()
            for target, names in HEADER_MAP.items():
                if key in names and v not in (None, ""):
                    canon[target] = str(v).strip()
        if canon.get("name"):
            out.append(canon)
    return out


def rows_to_raw(rows: list[dict], default_country: str | None, filename: str) -> list[RawCompany]:
    recs = []
    for i, r in enumerate(rows):
        country = (r.get("country") or default_country or "").upper()[:2] or None
        recs.append(RawCompany(
            name=r["name"], source_key="csv_import", source_url=f"upload://{filename}#row={i + 2}", source_record_id=f"{filename}:{i + 2}",
            country=country, city=r.get("city"), address=r.get("address"), phone=r.get("phone"), email=r.get("email"),
            website=r.get("website"), categories=[c.strip() for c in (r.get("categories") or "").split(",") if c.strip()],
            registration_number=r.get("registration_number"), vat_number=r.get("vat_number"),
            evidence_type="user_import", confidence=0.6,
        ))
    return recs


class CsvImportConnector(LeadSourceConnector):
    key = "csv_import"

    async def discover(self, query: DiscoveryQuery, context: ConnectorContext) -> DiscoveryResult:
        raise AccessRestrictedError("csv_import is driven by uploads (POST /api/companies/import), not by search queries")
